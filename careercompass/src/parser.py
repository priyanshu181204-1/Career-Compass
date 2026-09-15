"""
CareerCompass - Resume Parsing Engine
Supports: PDF, DOCX, and plain text.
Extracts: Contact info, education level, years of experience, and standardized technical skills.
"""

import os
import re
import json

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

def load_skills_taxonomy(kb_path=None):
    if kb_path is None:
        kb_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "knowledge_base", "skills_taxonomy.json")
    with open(kb_path, "r", encoding="utf-8") as f:
        taxonomy = json.load(f)
        
    skill_patterns = {}
    for cat, skills in taxonomy.items():
        for canon, meta in skills.items():
            names = [canon] + meta.get("synonyms", [])
            for n in names:
                escaped = re.escape(n)
                if n in ["C++", "C#", ".NET"]:
                    pattern = re.compile(rf"(?:^|[\s,;/]){escaped}(?:[\s,;/]|$)", re.IGNORECASE)
                else:
                    pattern = re.compile(rf"\b{escaped}\b", re.IGNORECASE)
                skill_patterns[canon] = skill_patterns.get(canon, []) + [pattern]
                
    extras = {
        "Git": [re.compile(r"\bGit\b", re.IGNORECASE)],
        "Agile": [re.compile(r"\bAgile\b|\bScrum\b", re.IGNORECASE)],
        "Linux": [re.compile(r"\bLinux\b|\bUnix\b|\bUbuntu\b", re.IGNORECASE)],
        "REST APIs": [re.compile(r"\bREST(?:ful)?\s+APIs?\b|\bREST\b", re.IGNORECASE)],
        "Unit Testing": [re.compile(r"\bUnit\s+Test(?:ing|s)?\b|\bPyTest\b|\bJUnit\b", re.IGNORECASE)],
        "CI/CD": [re.compile(r"\bCI/CD\b|\bContinuous\s+Integration\b", re.IGNORECASE)]
    }
    skill_patterns.update(extras)
    return skill_patterns

class ResumeParser:
    def __init__(self, kb_path=None):
        self.skill_patterns = load_skills_taxonomy(kb_path)
        
    def parse_file(self, file_path_or_buffer, filename=""):
        """Extract text from PDF, DOCX, or raw text."""
        ext = ""
        if isinstance(file_path_or_buffer, str):
            ext = os.path.splitext(file_path_or_buffer)[1].lower()
            if ext == ".pdf":
                text = self._extract_pdf(file_path_or_buffer)
            elif ext in [".docx", ".doc"]:
                text = self._extract_docx(file_path_or_buffer)
            else:
                with open(file_path_or_buffer, "r", encoding="utf-8", errors="ignore") as f:
                    text = f.read()
        else:
            ext = os.path.splitext(filename)[1].lower()
            if ext == ".pdf":
                text = self._extract_pdf_buffer(file_path_or_buffer)
            elif ext in [".docx", ".doc"]:
                text = self._extract_docx_buffer(file_path_or_buffer)
            else:
                text = file_path_or_buffer.read().decode("utf-8", errors="ignore")
                
        return self.parse_text(text)

    def _extract_pdf(self, path):
        text = ""
        if PdfReader:
            reader = PdfReader(path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        return text

    def _extract_pdf_buffer(self, buffer):
        text = ""
        if PdfReader:
            reader = PdfReader(buffer)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + "\n"
        return text

    def _extract_docx(self, path):
        text = ""
        if docx:
            doc = docx.Document(path)
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + "\n"
        return text

    def _extract_docx_buffer(self, buffer):
        text = ""
        if docx:
            doc = docx.Document(buffer)
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + "\n"
        return text

    def parse_text(self, text):
        skills_found = self.extract_skills(text)
        exp_years = self.extract_experience_years(text)
        edu_level, degree = self.extract_education(text)
        email, phone = self.extract_contact_info(text)
        
        return {
            "raw_text": text,
            "skills": skills_found,
            "years_experience": exp_years,
            "education_level": edu_level,
            "degree": degree,
            "email": email,
            "phone": phone
        }

    def extract_skills(self, text):
        found = set()
        for canon_skill, patterns in self.skill_patterns.items():
            for p in patterns:
                if p.search(text):
                    found.add(canon_skill)
                    break
        return sorted(list(found))

    def extract_experience_years(self, text):
        m = re.search(r"(\d{1,2})\+?\s*(?:years?|yrs?)(?:\s+of)?\s+experience", text, re.IGNORECASE)
        if m:
            return min(int(m.group(1)), 35)
            
        ranges = re.findall(r"\b(20\d\d|19\d\d)\s*[-–—]\s*(20\d\d|Present|Current)\b", text, re.IGNORECASE)
        total_yrs = 0
        for start, end in ranges:
            s_yr = int(start)
            e_yr = 2026 if end.lower() in ["present", "current"] else int(end)
            diff = max(0, e_yr - s_yr)
            total_yrs += diff
        if total_yrs > 0:
            return min(total_yrs, 35)
            
        return 2

    def extract_education(self, text):
        lower = text.lower()
        if re.search(r"\bph\.?d\b|\bdoctorate\b", lower):
            return "Doctorate", "Ph.D. in Computer Science or Related"
        if re.search(r"\bm\.?s\b|\bmaster(?:'s)?\b|\bm\.tech\b|\bmba\b", lower):
            return "Master", "Master of Science in Computer Science / Engineering"
        if re.search(r"\bb\.?s\b|\bbachelor(?:'s)?\b|\bb\.tech\b|\bb\.e\b", lower):
            return "Bachelor", "Bachelor of Science in Computer Science / Engineering"
        return "Bachelor", "Bachelor's Degree"

    def extract_contact_info(self, text):
        email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
        email = email_match.group(0) if email_match else ""
        phone_match = re.search(r"(?:\+?1\s*[-.]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
        phone = phone_match.group(0) if phone_match else ""
        return email, phone

if __name__ == "__main__":
    parser = ResumeParser()
    sample_text = """
    Alex Chen
    alex.chen@synthetic-talent.org | (555) 234-5678
    
    Professional Summary
    Senior Machine Learning Engineer with 6 years of experience building deep learning systems with PyTorch, Docker, and Kubernetes.
    
    Education
    M.S. in Computer Science, Carnegie Mellon University (2020)
    
    Technical Skills
    Python, PyTorch, Scikit-Learn, Docker, Kubernetes, AWS, SQL, CI/CD
    """
    parsed = parser.parse_text(sample_text)
    print("Parsed output test:")
    print("Skills:", parsed["skills"])
    print("Experience:", parsed["years_experience"])
    print("Education:", parsed["education_level"], f"({parsed['degree']})")
    print("Contact:", parsed["email"], parsed["phone"])
